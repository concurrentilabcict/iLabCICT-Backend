from django.db.models import Q,Subquery
from django.db import transaction
from openpyxl import load_workbook
from .models import Computer
from api.room.models import Room
from api.computer.models import Computer
from rest_framework.exceptions import ValidationError
from api.audit_logs.services import AuditLogsService
from channels.layers import get_channel_layer
from asgiref.sync import async_to_sync 
from api.ticket.services import TicketService
from django.db.models import Prefetch
from api.common.utils.entity_code import generate_entity_code
from api.common.utils.local_code import generate_pc_local_code
from django.utils import timezone
class ComputerService:

    @staticmethod
    def get_all_archived():
        return Computer.objects.filter(is_archived=True)

    @staticmethod
    @transaction.atomic
    def archive_computer(computer_id, request):
        from api.ticket.models import Ticket
        from api.maintenance_history.models import MaintenanceHistory
        from api.repair_log.models import RepairLog
        
        computer = (
            Computer.objects
            .select_for_update()
            .get(pk=computer_id)
        )

        room = computer.room

        if computer.is_archived:
            return computer

        computer.is_archived=True
        computer.save(update_fields=['is_archived'])

        ticket_ids = list(
                    Ticket.objects
                    .filter(computer_id=computer_id)
                    .values_list("id", flat=True)
                )
        
        repair_log_ids = list(
            RepairLog.objects
            .filter(ticket_id__in=ticket_ids)
            .values_list("id", flat=True)
        )

        maintenance_history_ids = list(
            MaintenanceHistory.objects
            .filter(repair_log_id__in=repair_log_ids)
            .values_list("id", flat=True)
        )

        MaintenanceHistory.objects.filter(
            id__in=maintenance_history_ids
        ).update(is_archived=True)

        RepairLog.objects.filter(
            id__in=repair_log_ids
        ).update(is_archived=True)

        Ticket.objects.filter(
            id__in=ticket_ids
        ).update(is_archived=True)

        AuditLogsService.log(
            request=request,
            performed_by=request.user,
            action_title="Archived computer",
            action_summary=(
                f"{request.user.get_full_name()} "
                f"archived a computer"
            ),
            metadata={
                "computer_id": computer_id,
                "affected_ticket_ids": ticket_ids,
                "affected_repair_log_ids": repair_log_ids,
                "affected_maintenance_history_ids": maintenance_history_ids,
            }
        )

        def broadcast():
            channel_layer = get_channel_layer()
            async_to_sync(channel_layer.group_send)(
                f'room_{room.id}',
                {
                    'type': 'computer_archived',
                    'computer_id': computer.id
                }
            )

        transaction.on_commit(broadcast)

    @staticmethod
    @transaction.atomic
    def unarchive_computer(computer_id, request):
        from api.ticket.models import Ticket
        from api.maintenance_history.models import MaintenanceHistory
        from api.repair_log.models import RepairLog

        computer = (
            Computer.objects
            .select_for_update()
            .get(pk=computer_id)
        )


        if not computer.is_archived:
            return computer

        computer.is_archived = False
        computer.save(update_fields=["is_archived"])

        room = computer.room

        ticket_ids = list(
            Ticket.objects
            .filter(computer_id=computer_id)
            .values_list("id", flat=True)
        )

        repair_log_ids = list(
            RepairLog.objects
            .filter(ticket_id__in=ticket_ids)
            .values_list("id", flat=True)
        )

        maintenance_history_ids = list(
            MaintenanceHistory.objects
            .filter(repair_log_id__in=repair_log_ids)
            .values_list("id", flat=True)
        )

        MaintenanceHistory.objects.filter(
            id__in=maintenance_history_ids
        ).update(is_archived=False)

        RepairLog.objects.filter(
            id__in=repair_log_ids
        ).update(is_archived=False)

        Ticket.objects.filter(
            id__in=ticket_ids
        ).update(is_archived=False)

        AuditLogsService.log(
            request=request,
            performed_by=request.user,
            action_title="Unarchived computer",
            action_summary=(
                f"{request.user.get_full_name()} "
                f"unarchived a computer"
            ),
            metadata={
                "computer_id": computer_id,
                "affected_ticket_ids": ticket_ids,
                "affected_repair_log_ids": repair_log_ids,
                "affected_maintenance_history_ids": maintenance_history_ids,
            }
        )

        def broadcast():
            from api.computer.serializers import ComputerDefaultSerializer
            channel_layer = get_channel_layer()
            serializer = ComputerDefaultSerializer(computer)
            async_to_sync(channel_layer.group_send)(
                f'room_{room.id}',
                {
                    'type': 'computer_unarchived',
                    'computer': serializer.data
                }
            )

        transaction.on_commit(broadcast)

    @staticmethod
    def get_computer_with_mainentance_history(include=None, computer_code=None):
        ticket_queryset = TicketService.get_tickets_per_computer()
        
        queryset = (Computer.objects
            .select_related('room')
            .prefetch_related(
                Prefetch(
                    'tickets',
                    queryset=ticket_queryset,
                    to_attr='assigned_tickets'
                )
            )
            )

        if "maintenance-history" in include.split(","):
            queryset = queryset.prefetch_related('maintenance_history', 
                                                 'maintenance_history__repair_log',
                                                 'maintenance_history__repair_log__ticket',
                                                 'maintenance_history__technician')

        return queryset

    #new method
    @staticmethod
    def get_all(filters):
        queryset = Computer.objects.select_related('room')

        

        ComputerService.validate_filters(filters)

        queryset = ComputerService.filter_per_computer_code(queryset, filters)
        queryset = ComputerService.filter_active(queryset, filters)
        queryset = ComputerService.filter_all_peripherals(queryset, filters)
        queryset = ComputerService.filter_peripheral_status(queryset, filters)


        include = filters.include or ""

        if "maintenance-history" in include.split(","):
            queryset = queryset.prefetch_related("maintenance_history")


        return queryset
    
    def filter_per_computer_code(queryset, filters):
        computer_code = filters.get('computer-code')
        if filters.get('computer-code') is not None:
            queryset = queryset.filter(computer_code=computer_code)

        return queryset
    
    @staticmethod
    def filter_active(queryset, filters):
        if filters.get('active') == 'true':
            queryset = queryset.filter(computer_status=Computer.ComputerStatus.ACTIVE)
        
        return queryset
    
    @staticmethod
    def filter_all_peripherals(queryset, filters):
        all_peripheral_status = filters.get('peripherals')
        
        if all_peripheral_status == 'none':
            queryset = queryset.filter(
                mouse_status=Computer.PeripheralStatus.NONE,
                keyboard_status=Computer.PeripheralStatus.NONE,
                monitor_status=Computer.PeripheralStatus.NONE,
                ups_status=Computer.PeripheralStatus.NONE
            )

        elif all_peripheral_status == 'all':
            queryset = queryset.filter(
                mouse_status=Computer.PeripheralStatus.ACTIVE,
                keyboard_status=Computer.PeripheralStatus.ACTIVE,
                monitor_status=Computer.PeripheralStatus.ACTIVE,
                ups_status=Computer.PeripheralStatus.ACTIVE
            )

        return queryset
    
    @staticmethod
    def filter_peripheral_status(queryset, filters):
        peripheral = filters.get('peripheral-type')
        status = filters.get('status')

        if peripheral and status:
            queryset = queryset.filter(
                **{f"{peripheral}_status": status}
            )

        return queryset


    @staticmethod
    def validate_filters(filters):
        allowed_peripheral_types=[
            'keyboard',
            'ups',
            'monitor',
            'mouse'
        ]

        allowed_statuses = Computer.ComputerStatus.values

        peripheral = filters.get('peripheral-type')
        status = filters.get('status')

        if peripheral and peripheral not in allowed_peripheral_types:
            raise ValidationError('Invalid peripheral type')
        
        if status and status not in allowed_statuses:
            raise ValidationError('Invalid peripheral status')


    @staticmethod
    def create_computers(serializer, request):
        computers = serializer.save()

        ComputerService.broadcast_computer_created(
            computers=computers,
            event_type='computer_created'
        )

        AuditLogsService.log(
            request=request,
            performed_by=request.user,
            action_title='Computers created',
            action_summary=f"{request.user.get_full_name()} created {len(computers)} computer(s).",
            metadata={
                "computer_ids": [computer.id for computer in computers],
                "quantity": len(computers),
            }
        )

        return computers

    @staticmethod
    def update_computer(serializer, request):
        from api.computer.serializers import ComputerDefaultSerializer
        computer = serializer.instance

        old_values = {}

        for field, new_value in serializer.validated_data.items():
            old_value = getattr(computer, field)

            if old_value != new_value:
                old_values[field] = old_value

        computer = serializer.save()

        changes = {}

        for field in old_values:
            changes[field] = {
                'old': old_values[field],
                'new': getattr(computer, field)
            }

        ComputerService.broadcast_computer_updated(
            computer=computer,
            event_type='computer_updated'
        )

        AuditLogsService.log(
            request=request,
            performed_by=request.user,
            action_title='Computer updated',
            action_summary=f"{request.user.get_full_name()} updated computer '{computer.computer_code}.'",
            metadata={
                'computer_id': computer.id,
                'changes': changes
            }
        )

        return computer

    @staticmethod
    def broadcast_computer_created(computers, event_type):
        from api.computer.serializers import ComputerDefaultSerializer

        channel_layer = get_channel_layer()

        room_id = computers[0].room_id

        serialized_computers = ComputerDefaultSerializer(
            computers,
            many=True
        ).data

        async_to_sync(channel_layer.group_send)(
            f'room_{room_id}',
            {
                'type': event_type,
                'computer': serialized_computers
            }
        )

    @staticmethod
    def broadcast_computer_updated(computer, event_type):
        from api.computer.serializers import ComputerDefaultSerializer
        channel_layer = get_channel_layer()

        room_id = computer.room_id

        serialized_computer = ComputerDefaultSerializer(computer).data

        async_to_sync(channel_layer.group_send)(
            f'room_{room_id}',
            {
                'type': event_type,
                'computer': serialized_computer
            }
        )


#---------------------------------------------old method-----------------------------------------------------------
    @staticmethod
    def get_all_active():
        return Computer.objects.filter(computer_status=Computer.ComputerStatus.ACTIVE)
    
    @staticmethod
    def get_all_with_active_peripherals():
        return Computer.objects.filter(
            computer_status=Computer.ComputerStatus.ACTIVE,
            mouse_status=Computer.PeripheralStatus.ACTIVE,
            monitor_status=Computer.PeripheralStatus.ACTIVE,
            keyboard_status=Computer.PeripheralStatus.ACTIVE,
            ups_status=Computer.PeripheralStatus.ACTIVE
            )
    
    @staticmethod
    def get_all_active_with_peripheral(filters):
        queryset = Computer.objects.filter(
            Q(computer_status=Computer.ComputerStatus.ACTIVE) &
            (
                ~Q(mouse_status=Computer.PeripheralStatus.NONE) &
                ~Q(keyboard_status=Computer.PeripheralStatus.NONE) &
                ~Q(monitor_status=Computer.PeripheralStatus.NONE) &
                ~Q(ups_status=Computer.PeripheralStatus.NONE)
            )
            )

        peripheral_type = filters.get("type")
        status = filters.get("status")

        if peripheral_type:
            queryset = queryset.filter(**{
                f"{peripheral_type}_status": f"{status}"
                })
        
        return queryset
    
    @staticmethod
    def get_all_active_no_peripherals():
        return Computer.objects.filter(
            Q(mouse_status=Computer.PeripheralStatus.NONE) &
            Q(keyboard_status=Computer.PeripheralStatus.NONE) &
            Q(monitor_status=Computer.PeripheralStatus.NONE) &
            Q(ups_status=Computer.PeripheralStatus.NONE)
        )   

    REQUIRED_COLUMNS = {
        "room",
        "model",
        "processor",
        "ram",
        "storage",
        "operating_system",
    }
    
    REQUIRED_COLUMNS = {
        "operating_system",
        "gpu",
        "cpu",
        "ram_size_installed",
        "disk_size_installed",
        "build_version",
        "computer_status",
        "monitor_status",
        "mouse_status",
        "keyboard_status",
        "ups_status",
        "motherboard",
    }

    @classmethod
    def import_file(cls, excel_file, room_id):

        # ---------------------------------------------------------
        # 1. Validate file
        # ---------------------------------------------------------

        if not excel_file:
            raise ValueError(
                "No Excel file was provided."
            )

        if not excel_file.name.lower().endswith(".xlsx"):
            raise ValueError(
                "Only .xlsx files are supported."
            )

        # ---------------------------------------------------------
        # 2. Get room
        # ---------------------------------------------------------

        try:
            room = Room.objects.get(
                pk=room_id
            )

        except Room.DoesNotExist:
            raise ValueError(
                f"Room with ID {room_id} does not exist."
            )

        # ---------------------------------------------------------
        # 3. Load workbook
        # ---------------------------------------------------------

        try:
            workbook = load_workbook(
                excel_file,
                read_only=True,
                data_only=True,
            )

            worksheet = workbook.active

        except Exception:
            raise ValueError(
                "Unable to read the Excel file."
            )

        # ---------------------------------------------------------
        # 4. Read headers
        # ---------------------------------------------------------

        rows = worksheet.iter_rows(
            values_only=True
        )

        try:
            raw_headers = next(rows)

        except StopIteration:
            raise ValueError(
                "The Excel file is empty."
            )

        headers = [
            str(header).strip().lower()
            if header is not None
            else ""
            for header in raw_headers
        ]

        # Remove empty headers
        headers = [
            header
            for header in headers
            if header
        ]

        # ---------------------------------------------------------
        # 5. Validate headers
        # ---------------------------------------------------------

        missing_columns = (
            cls.REQUIRED_COLUMNS - set(headers)
        )

        if missing_columns:
            raise ValueError(
                {
                    "message": "Missing required columns.",
                    "missing_columns": sorted(
                        missing_columns
                    ),
                }
            )

        # ---------------------------------------------------------
        # 6. Parse rows
        # ---------------------------------------------------------

        computers = []
        errors = []

        for row_number, row in enumerate(
            rows,
            start=2,
        ):

            # Skip completely empty rows
            if all(
                value is None
                or str(value).strip() == ""
                for value in row
            ):
                continue

            data = dict(
                zip(headers, row)
            )

            row_errors = []

            # -----------------------------------------------------
            # Required text fields
            # -----------------------------------------------------

            operating_system = cls.clean_value(
                data.get("operating_system")
            )

            gpu = cls.clean_value(
                data.get("gpu")
            )

            cpu = cls.clean_value(
                data.get("cpu")
            )

            build_version = cls.clean_value(
                data.get("build_version")
            )

            computer_status = cls.clean_value(
                data.get("computer_status")
            )

            monitor_status = cls.clean_value(
                data.get("monitor_status")
            )

            mouse_status = cls.clean_value(
                data.get("mouse_status")
            )

            keyboard_status = cls.clean_value(
                data.get("keyboard_status")
            )

            ups_status = cls.clean_value(
                data.get("ups_status")
            )

            motherboard = cls.clean_value(
                data.get("motherboard")
            )

            # -----------------------------------------------------
            # Required integer fields
            # -----------------------------------------------------

            ram_size_installed = cls.clean_integer(
                data.get("ram_size_installed")
            )

            disk_size_installed = cls.clean_integer(
                data.get("disk_size_installed")
            )

            # -----------------------------------------------------
            # Validate text fields
            # -----------------------------------------------------

            required_text_fields = {
                "operating_system": operating_system,
                "gpu": gpu,
                "cpu": cpu,
                "build_version": build_version,
                "computer_status": computer_status,
                "monitor_status": monitor_status,
                "mouse_status": mouse_status,
                "keyboard_status": keyboard_status,
                "ups_status": ups_status,
                "motherboard": motherboard,
            }

            for field_name, value in (
                required_text_fields.items()
            ):
                if not value:
                    row_errors.append(
                        f"{field_name} is required."
                    )

            # -----------------------------------------------------
            # Validate integer fields
            # -----------------------------------------------------

            if ram_size_installed is None:
                row_errors.append(
                    "ram_size_installed must be an integer."
                )

            if disk_size_installed is None:
                row_errors.append(
                    "disk_size_installed must be an integer."
                )

            # -----------------------------------------------------
            # Row validation failed
            # -----------------------------------------------------

            if row_errors:
                errors.append(
                    {
                        "row": row_number,
                        "errors": row_errors,
                    }
                )

                continue

            # -----------------------------------------------------
            # Create unsaved Computer
            # -----------------------------------------------------

            computer = Computer(
                room=room,

                # Generated by this import service
                computer_code=None,
                computer_number=None,

                operating_system=operating_system,
                gpu=gpu,
                cpu=cpu,
                ram_size_installed=ram_size_installed,
                disk_size_installed=disk_size_installed,
                build_version=build_version,
                computer_status=computer_status,
                monitor_status=monitor_status,
                mouse_status=mouse_status,
                keyboard_status=keyboard_status,
                ups_status=ups_status,
                motherboard=motherboard,

                # System/default field
                is_archived=False,
            )

            computers.append(computer)

        # ---------------------------------------------------------
        # 7. Stop if validation errors exist
        # ---------------------------------------------------------

        if errors:
            raise ValueError(
                {
                    "message": (
                        "Import failed because "
                        "some rows are invalid."
                    ),
                    "errors": errors,
                }
            )

        if not computers:
            raise ValueError(
                "No valid computer records were found."
            )

        # ---------------------------------------------------------
        # 8. Generate identifiers + bulk insert
        # ---------------------------------------------------------

        try:

            with transaction.atomic():

                cls.generate_computer_numbers(
                    computers,
                    room,
                )

                cls.generate_computer_codes(
                    computers
                )

                Computer.objects.bulk_create(
                    computers
                )

        except Exception as e:

            raise ValueError(
                f"Failed to import computers: {str(e)}"
            )

        return {
            "created": len(computers),
        }

    # =============================================================
    # COMPUTER NUMBER GENERATION
    # =============================================================

    @staticmethod
    def generate_computer_numbers(
        computers,
        room,
    ):
        """
        Generate PC-XX numbers for the specified room.
        """

        existing_numbers = set(
            Computer.objects
            .filter(room=room)
            .values_list(
                "computer_number",
                flat=True,
            )
        )

        next_number = 1

        for computer in computers:

            while (
                f"PC-{next_number:02d}"
                in existing_numbers
            ):
                next_number += 1

            computer.computer_number = (
                f"PC-{next_number:02d}"
            )

            existing_numbers.add(
                computer.computer_number
            )

            next_number += 1

    # =============================================================
    # COMPUTER CODE GENERATION
    # =============================================================

    @staticmethod
    def generate_computer_codes(computers):

        if not computers:
            return

        current_year = timezone.now().year

        last_obj = (
            Computer.objects
            .filter(
                computer_code__startswith=f"PC{current_year}"
            )
            .order_by("-computer_code")
            .first()
        )

        if last_obj:
            last_number = int(
                last_obj.computer_code[-5:]
            )
        else:
            last_number = 0

        for index, computer in enumerate(
            computers,
            start=1,
        ):
            number = last_number + index

            computer.computer_code = (
                f"PC{current_year}{number:05d}"
            )

    # =============================================================
    # HELPERS
    # =============================================================

    @staticmethod
    def clean_value(value):

        if value is None:
            return None

        value = str(value).strip()

        return value if value else None

    @staticmethod
    def clean_integer(value):

        if value is None:
            return None

        if isinstance(value, bool):
            return None

        try:
            return int(value)

        except (ValueError, TypeError):
            return None


